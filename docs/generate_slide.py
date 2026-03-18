"""Generate a single summary slide for tiny-trtllm."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# Background
bg = slide.background
fill = bg.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0x0D, 0x11, 0x17)

# --- Title ---
txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(14), Inches(1.2))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run1 = p.add_run()
run1.text = "tiny-"
run1.font.size = Pt(52)
run1.font.bold = True
run1.font.color.rgb = RGBColor(0xE6, 0xED, 0xF3)
run2 = p.add_run()
run2.text = "trtllm"
run2.font.size = Pt(52)
run2.font.bold = True
run2.font.color.rgb = RGBColor(0x58, 0xA6, 0xFF)

# Subtitle
txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(14), Inches(0.6))
tf2 = txBox2.text_frame
p2 = tf2.paragraphs[0]
run_sub = p2.add_run()
run_sub.text = "Minimal reimplementation of TensorRT-LLM's PyTorch backend  |  ~3,700 LOC  |  206 tests  |  86% coverage  |  0 C++ files"
run_sub.font.size = Pt(18)
run_sub.font.color.rgb = RGBColor(0x8B, 0x94, 0x9E)

# --- Divider line ---
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.1), Inches(14.4), Pt(2))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0x30, 0x36, 0x3D)
shape.line.fill.background()

# --- Left column: 16 Core Differentiators ---
left_x = Inches(0.8)

txBox3 = slide.shapes.add_textbox(left_x, Inches(2.4), Inches(7), Inches(0.5))
tf3 = txBox3.text_frame
p3 = tf3.paragraphs[0]
run_h = p3.add_run()
run_h.text = "16 Core Architectural Differentiators Preserved"
run_h.font.size = Pt(22)
run_h.font.bold = True
run_h.font.color.rgb = RGBColor(0x7E, 0xE7, 0x87)

features = [
    ("Engine", [
        "Two-tier scheduling (Capacity + MicroBatch)",
        "GUARANTEED_NO_EVICT / MAX_UTILIZATION policies",
        "EQUAL_PROGRESS / FCFS chunking policies",
        "Iteration-based executor loop",
        "Phase-segmented ScheduledRequests (4 lists)",
        "Async two-phase sampling",
        "Grouped strategy sampling",
        "Overlap executor (CPU/GPU ping-pong)",
        "Await-responses (threading.Condition)",
    ]),
    ("Layers", [
        "Plan/execute attention (plan → run)",
        "Hybrid attention backend dispatcher",
        "TRT-LLM thop.attention() kernel (via pip)",
        "MoE (gate → topk → fused dispatch)",
    ]),
    ("Infra", [
        "Pydantic config hierarchy",
        "ResourceManager lifecycle (prepare → update → free)",
        "Multi-GPU TP with NCCL",
    ]),
]

y = Inches(3.0)
for category, items in features:
    txBox_cat = slide.shapes.add_textbox(left_x, y, Inches(7), Inches(0.3))
    tf_cat = txBox_cat.text_frame
    p_cat = tf_cat.paragraphs[0]
    run_cat = p_cat.add_run()
    run_cat.text = category
    run_cat.font.size = Pt(14)
    run_cat.font.bold = True
    run_cat.font.color.rgb = RGBColor(0x58, 0xA6, 0xFF)
    y += Inches(0.3)

    for item in items:
        txBox_item = slide.shapes.add_textbox(left_x + Inches(0.2), y, Inches(6.8), Inches(0.25))
        tf_item = txBox_item.text_frame
        p_item = tf_item.paragraphs[0]
        run_item = p_item.add_run()
        run_item.text = f"  {item}"
        run_item.font.size = Pt(12)
        run_item.font.color.rgb = RGBColor(0x8B, 0x94, 0x9E)
        y += Inches(0.25)
    y += Inches(0.1)

# --- Right column: Architecture + Stats ---
right_x = Inches(8.5)

txBox4 = slide.shapes.add_textbox(right_x, Inches(2.4), Inches(7), Inches(0.5))
tf4 = txBox4.text_frame
p4 = tf4.paragraphs[0]
run_h2 = p4.add_run()
run_h2.text = "Architecture"
run_h2.font.size = Pt(22)
run_h2.font.bold = True
run_h2.font.color.rgb = RGBColor(0xD2, 0xA8, 0xFF)

# Architecture diagram as text
arch_text = """tinytrtllm/
  config.py          Pydantic config
  llm.py             Top-level API
  engine/
    scheduler.py     Two-tier scheduling
    executor.py      PyExecutor + Overlap
    sampler.py       Async two-phase
    block_manager.py Paged KV + prefix cache
    model_engine.py  Forward + CUDA graphs
  layers/
    attention*.py    5 backends + hybrid
    linear.py        TP-aware parallel
    moe.py           Gate + TopK + fused
  models/
    llama.py         LlamaForCausalLM
    qwen3.py         Qwen3ForCausalLM
    qwen3_moe.py     Qwen3MoEForCausalLM
  serve/server.py    FastAPI OpenAI-compat"""

txBox_arch = slide.shapes.add_textbox(right_x, Inches(3.0), Inches(7), Inches(3.5))
tf_arch = txBox_arch.text_frame
tf_arch.word_wrap = True
for line in arch_text.strip().split("\n"):
    p_line = tf_arch.add_paragraph()
    run_line = p_line.add_run()
    run_line.text = line
    run_line.font.size = Pt(11)
    run_line.font.name = "Courier New"
    run_line.font.color.rgb = RGBColor(0x8B, 0x94, 0x9E)

# Stats boxes
stats = [
    ("3,654", "Production LOC", "7EE787"),
    ("2,578", "Test LOC", "58A6FF"),
    ("86%", "Coverage", "D2A8FF"),
    ("206", "Tests Passing", "FFA657"),
]

stat_y = Inches(6.8)
for i, (val, label, color) in enumerate(stats):
    sx = right_x + Inches(i * 1.8)
    # Value
    txBox_v = slide.shapes.add_textbox(sx, stat_y, Inches(1.6), Inches(0.5))
    tf_v = txBox_v.text_frame
    p_v = tf_v.paragraphs[0]
    p_v.alignment = PP_ALIGN.CENTER
    run_v = p_v.add_run()
    run_v.text = val
    run_v.font.size = Pt(28)
    run_v.font.bold = True
    run_v.font.color.rgb = RGBColor(int(color[:2], 16), int(color[2:4], 16), int(color[4:], 16))
    # Label
    txBox_l = slide.shapes.add_textbox(sx, stat_y + Inches(0.45), Inches(1.6), Inches(0.3))
    tf_l = txBox_l.text_frame
    p_l = tf_l.paragraphs[0]
    p_l.alignment = PP_ALIGN.CENTER
    run_l = p_l.add_run()
    run_l.text = label
    run_l.font.size = Pt(11)
    run_l.font.color.rgb = RGBColor(0x8B, 0x94, 0x9E)

# What's removed
txBox_rm = slide.shapes.add_textbox(right_x, Inches(7.8), Inches(7), Inches(0.8))
tf_rm = txBox_rm.text_frame
tf_rm.word_wrap = True
p_rm = tf_rm.paragraphs[0]
run_rm = p_rm.add_run()
run_rm.text = "Removed: TensorRT compilation, speculative decoding, disaggregated serving, pipeline parallelism, LoRA, quantization, beam search, multimodal, 25+ architectures, C++ bindings"
run_rm.font.size = Pt(10)
run_rm.font.italic = True
run_rm.font.color.rgb = RGBColor(0x6B, 0x74, 0x7E)

# Footer
txBox_f = slide.shapes.add_textbox(Inches(0.8), Inches(8.5), Inches(14), Inches(0.3))
tf_f = txBox_f.text_frame
p_f = tf_f.paragraphs[0]
run_f = p_f.add_run()
run_f.text = "github.com/venkywonka/tiny-trtllm  |  branch: dev/poc-1  |  Python + Triton only, zero C++"
run_f.font.size = Pt(12)
run_f.font.color.rgb = RGBColor(0x58, 0xA6, 0xFF)

prs.save("/home/ubuntu/tiny-trtllm/docs/tiny-trtllm-summary.pptx")
print("Saved: docs/tiny-trtllm-summary.pptx")
